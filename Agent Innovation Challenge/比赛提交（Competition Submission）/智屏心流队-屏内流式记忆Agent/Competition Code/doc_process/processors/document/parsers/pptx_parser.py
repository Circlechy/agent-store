# -*- coding: utf-8 -*-
# Copyright (c) huawei, Inc. and its affiliates.
import pptx

from doc_process.processors.document.parsers.doc_default_parser import DocumentDefaultParser
from doc_process.context.doc_schema import TextContent, TableContent
from doc_process.context.doc_schema import DocNode, Document
from doc_process.utils import logging
from doc_process.utils.check_utils import CheckUtils
from doc_process.utils.error_code import ProcessorException, ErrorCode
from doc_process.context.doc_constants import PARA_TYPE, COMB_TYPE, TABLE_TYPE, DEFAULT_DEGRADE_ENABLE

logger = logging.get_logger()


class PPTXParser(DocumentDefaultParser):
    """ PPTXParser """

    def __init__(self, **kwargs):
        super(PPTXParser, self).__init__(**kwargs)


    @classmethod
    def file_suffix(cls) -> str:
        """file suffix"""
        return ".pptx"

    @classmethod
    def convert_2d_list(cls, table):
        """ convert table to 2d list """
        if not table:
            return []
        max_col = max(len(t) for t in table)
        if max_col == 0:
            return []
        new_table = []
        for row in table:
            row = [cell.replace("\n", "<br>").replace("|", "&#124;")
                   if isinstance(cell, str) else ""
                   for cell in row]
            row = row + [""] * (max_col - len(row))
            new_table.append(row)
        return new_table


    @classmethod
    def convert_md_row(cls, row, max_col, pad_txt=""):
        """
        Convert 1D list to markdown row
        """
        row = [cell.replace("\n", "<br>").replace("|", "&#124;")
               if isinstance(cell, str) else pad_txt
               for cell in row]
        row = row + [pad_txt] * (max_col - len(row))
        return "|" + "|".join(row) + "|"


    @classmethod
    def get_shape_position(cls, shape):
        """ get shape position """
        position = (shape.top, shape.left)
        if shape.has_text_frame:
            position = (position[0] or shape.text_frame.margin_top, position[1] or shape.text_frame.margin_left)
        if None in position:
            position = (float("inf"), float("inf"))
        return position


    def title_scoring(self, feat):
        """
        Title scoring
        feat: [y, x]
        """
        return abs(feat[0]) * self.cfg.ppt_page_title_score_weight[0] + \
            abs(feat[1]) * self.cfg.ppt_page_title_score_weight[1]


    def read_from_path(self, path):
        """ read from file """
        presentation = pptx.Presentation(path)
        if self.degrade:
            self.check_large_file(presentation)
        nodes = self.slides_to_docnode(path.stem, presentation)
        doc_tree = Document.tree_from_nodes(nodes)
        return doc_tree, {"p_type": "ppt"}

    def check_large_file(self, presentation):
        """限制超大文件"""
        num_pages = len(presentation.slides)
        num_slide_shapes = 0

        if num_pages > self.cfg.max_page_limit:
            raise ProcessorException(ErrorCode.FILE_TOO_LARGE,
                                     "File has exceeded the maximum limit of {} pages. Current page count: {} "
                                     .format(self.cfg.max_page_limit, num_pages))

        for slide in presentation.slides:
            num_slide_shapes += len(slide.shapes)

            if num_slide_shapes > self.cfg.max_shapes_limit:
                raise ProcessorException(ErrorCode.FILE_TOO_LARGE,
                                         "File has exceeded the maximum limit of {} slide_shapes. "
                                         "Current slide_shapes count: {} "
                                         .format(self.cfg.max_shapes_limit, num_slide_shapes))

    def slides_to_docnode(self, file_title, presentation):
        """
        Build DocNode list from each page
        """
        root_node = DocNode(id="", chapter=[file_title], title_level=0, content=[], parent=None,
                            children=[])
        nodes = [root_node]

        for i, slide in enumerate(presentation.slides):
            slide_nodes = self.extract_shapes(slide.shapes, 0, i + 1)

            if slide_nodes:
                nodes.append(DocNode(id="", chapter=[file_title, ""], title_level=1, content=[], parent=None,
                                     children=[]))
                # Find title (the top left paragraph)
                slide_nodes.sort(key=lambda x: self.title_scoring(x[0]))
                self.add_node_content(nodes, slide_nodes)
        # Update DocNode.id
        for i, node in enumerate(nodes):
            node.id = i
        return nodes

    def add_node_content(self, nodes, slide_nodes):
        """Add node content"""
        if slide_nodes[0][1] == PARA_TYPE:
            nodes[-1].chapter[-1] = slide_nodes[0][2]
            slide_nodes.pop(0)
        # Content, text and table
        for row in slide_nodes:
            if row[1] == TABLE_TYPE:
                nodes[-1].content.append(TableContent(table=row[2]))
            else:
                nodes[-1].content.append(TextContent(content=row[2]))

    def extract_shapes(self, shapes, depth, slide_idx):
        """
        Extract content from shape list
        """
        slide_nodes = []
        for shape in shapes:
            self.extract_node_from_shape(depth, shape, slide_idx, slide_nodes)
        slide_nodes.sort(key=lambda x: x[0])
        return slide_nodes


    def extract_node_from_shape(self, depth, shape, slide_idx, slide_nodes):
        """ Extract node from pptx shape """
        position = self.get_shape_position(shape)
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                para_text = "".join([run.text for run in paragraph.runs])
                if not para_text:
                    continue
                slide_nodes.append([position, PARA_TYPE, para_text])
        elif shape.has_table:
            table = [
                [cell.text for cell in row.cells]
                for row in shape.table.rows
            ]
            table_text = self.convert_2d_list(table)
            slide_nodes.append([position, TABLE_TYPE, table_text])
        elif hasattr(shape, "shapes"):
            tmp = self.extract_shapes(shape.shapes, depth + 1, slide_idx)
            slide_nodes.extend([[position, COMB_TYPE, x[-1]] for x in tmp])


    def _parse_config(self, **kwargs):
        self.degrade = self.config.get("degrade", {}).get("enable", DEFAULT_DEGRADE_ENABLE)
        self.config = self.config.get("parser", {}).get("pptx", {})

    def _check_cfg(self, cfg):
        """ check config """
        CheckUtils.check_iter(cfg.ppt_page_title_score_weight, length=2,
                              name_="parser.pptx.ppt_page_title_score_weight")
        for weight in cfg.ppt_page_title_score_weight:
            CheckUtils.check_types(weight, [float, int], "parser.pptx.ppt_page_title_score_weight")
            CheckUtils.check_range(weight, min_close=0.0, max_close=1.0)
        CheckUtils.check_type(cfg.max_page_limit, int, "parser.pptx.max_page_limit")
        CheckUtils.check_type(cfg.max_shapes_limit, int, "parser.pptx.max_shapes_limit")
        CheckUtils.check_types(cfg.file_size_limit, [int, float], "parser.pptx.file_size_limit")
