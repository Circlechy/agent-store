# -*- coding: utf-8 -*-
# Copyright (c) huawei, Inc. and its affiliates.

from openpyxl import load_workbook

from doc_process.processors.document.parsers.doc_default_parser import DocumentDefaultParser
from doc_process.utils import logging
from doc_process.context.doc_schema import DocNode, Document, TextContent
from doc_process.utils.check_utils import CheckUtils
from doc_process.context.doc_constants import DEFAULT_DEGRADE_ENABLE
from doc_process.utils.error_code import ProcessorException, ErrorCode

logger = logging.get_logger()


class ExcelParser(DocumentDefaultParser):
    """ ExcelParser """

    def __init__(self, **kwargs):
        super(ExcelParser, self).__init__(**kwargs)

    @classmethod
    def file_suffix(cls) -> str:
        """file suffix"""
        return ".xlsx"

    @classmethod
    def build_excel_nodes(cls, text, fnm):
        """ build node list from text """
        nodes = [DocNode(id="", chapter=[fnm], title_level=0, content=[], parent=None, children=[])]
        # 只读第一个sheet页
        sheet = text.sheetnames[0]
        work_sheet = text[sheet]
        rows = list(work_sheet.rows)

        if rows:
            for idx, r in enumerate(list(rows[1:])):
                question = str(r[0].value) if r[0].value else ''
                answer = str(r[1].value) if r[1].value else ''

                if question.strip() and answer.strip():
                    nodes.append(DocNode(id="", chapter=[fnm, answer], title_level=1,
                                         content=[TextContent(content=question)], parent=None, children=[]))
                elif question.strip() or answer.strip():
                    raise ProcessorException(ErrorCode.UNSUPPORTED_FILE_FORMAT,
                                             "File has empty cell in row: {}".format(idx+2))
        # udpate DocNode.id
        for i, node in enumerate(nodes):
            node.id = i
        return nodes

    def check_illegal_file(self, text):
        """检测不合法文件"""
        work_sheet = text[text.sheetnames[0]]
        rows = list(work_sheet.rows)
        len_cols = work_sheet.max_column
        merged_range = work_sheet.merged_cells.ranges

        if len_cols != self.cfg.col_limit:
            raise ProcessorException(ErrorCode.UNSUPPORTED_FILE_FORMAT,
                                     "File does not meet the column requirement."
                                     " Current columns count: {}, Only support columns: {}."
                                     .format(len_cols, self.cfg.col_limit))

        if merged_range:
            raise ProcessorException(ErrorCode.UNSUPPORTED_FILE_FORMAT,
                                     "File has merged cells. Merged range: {} "
                                     .format(merged_range))

        if rows:
            self.check_cell(rows)

    def check_cell(self, rows):
        """检测单元格"""
        for row in list(rows):

            if len(str(row[0].value)) > self.cfg.max_1st_col_char:
                raise ProcessorException(ErrorCode.FILE_TOO_LARGE,
                                         "File has exceeded the maximum limit of {} characters "
                                         "in first column cell. "
                                         "Current cell character count: {} "
                                         .format(self.cfg.max_1st_col_char, len(str(row[0].value))))

            if self.degrade and len(str(row[1].value)) > self.cfg.max_2nd_col_char:
                raise ProcessorException(ErrorCode.FILE_TOO_LARGE,
                                         "File has exceeded the maximum limit of {} characters "
                                         "in second column. "
                                         "Current cell character count: {} "
                                         .format(self.cfg.max_2nd_col_char, len(str(row[1].value))))

    def check_large_file(self, text_result):
        """限制超大文件"""
        sheet = text_result.sheetnames[0]
        work_sheet = text_result[sheet]
        rows = list(work_sheet.rows)
        if len(rows) > self.cfg.max_row_limit:
            raise ProcessorException(ErrorCode.FILE_TOO_LARGE,
                                     "File has exceeded the maximum limit of {} rows. "
                                     "Current rows count: {} "
                                     .format(self.cfg.max_row_limit, len(rows)))

    def read_from_path(self, path):
        """ read from file """
        text_result = load_workbook(path)
        if self.degrade:
            self.check_large_file(text_result)
        self.check_illegal_file(text_result)
        doc_nodes = self.build_excel_nodes(text_result, path.stem)
        doc_tree = Document.tree_from_nodes(doc_nodes)
        return doc_tree, {"p_type": "faq"}

    def _parse_config(self, **kwargs):
        self.degrade = self.config.get("degrade", {}).get("enable", DEFAULT_DEGRADE_ENABLE)
        self.config = self.config.get("parser", {}).get("xlsx", {})

    def _check_cfg(self, cfg):
        """ check config """
        CheckUtils.check_type(cfg.max_row_limit, int, "parser.xlsx.max_row_limit")
        CheckUtils.check_types(cfg.file_size_limit, [int, float], "parser.xlsx.file_size_limit")
