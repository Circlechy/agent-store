from pptx import Presentation
from Quartz import (
    CGWindowListCopyWindowInfo,
    kCGWindowListOptionOnScreenOnly,
    kCGNullWindowID,
)

# 使用 AppKit 获取屏幕尺寸（线程安全，无需 Tkinter）
try:
    from AppKit import NSScreen

    def get_screen_size():
        main_screen = NSScreen.mainScreen()
        screen_frame = main_screen.frame()
        return int(screen_frame.size.width), int(screen_frame.size.height)

except ImportError:
    # 如果 AppKit 不可用，使用固定值（兜底）
    def get_screen_size():
        return 1920, 1080  # 默认分辨率


class PPTDataParser:
    """PPT解析层（兼容全屏/普通模式）"""

    def __init__(self, offset_y=-20):
        self.offset_y = offset_y
        self.elem_texts = {}
        self.elem_coords = {}
        self.ppt_win = {}

    def parse_ppt(self, ppt_path):
        print(f"[Parser] 开始解析PPT：{ppt_path}")
        prs = Presentation(ppt_path)
        slide = prs.slides[0]
        print(f"[Parser] PPT第1页，形状数量：{len(slide.shapes)}")

        # 获取PPT尺寸
        try:
            slide_width = slide.layout.slide_master.slide_width
            slide_height = slide.layout.slide_master.slide_height
        except:
            slide_width = 12192000
            slide_height = 6858000
        sw = float(slide_width)
        sh = float(slide_height)

        # 识别PPT窗口
        window_list = CGWindowListCopyWindowInfo(
            kCGWindowListOptionOnScreenOnly, kCGNullWindowID
        )
        ppt_keywords = ["powerpoint", "keynote", "ppt", "演示文稿", "幻灯片", "wps office"]
        ppt_win = None

        # 屏幕尺寸 - 使用 AppKit（线程安全，无 Tkinter 依赖）
        screen_width, screen_height = get_screen_size()
        fullscreen_threshold = 0.9

        # 遍历窗口
        for win in window_list:
            owner = win.get("kCGWindowOwnerName", "").lower()
            is_onscreen = win.get("kCGWindowIsOnscreen", False)
            if is_onscreen and any(k in owner for k in ppt_keywords):
                bounds = win["kCGWindowBounds"]
                win_width = float(bounds["Width"])
                win_height = float(bounds["Height"])

                is_fullscreen = (win_width >= screen_width * fullscreen_threshold) and (
                    win_height >= screen_height * fullscreen_threshold
                )

                if is_fullscreen:
                    ppt_win = {
                        "x": 0.0,
                        "y": 0.0,
                        "width": screen_width,
                        "height": screen_height,
                        "is_fullscreen": True,
                    }
                    print(
                        f"[Parser] 检测到PPT全屏模式，强制使用屏幕坐标：{screen_width}x{screen_height}"
                    )
                    break
                else:
                    ppt_win = {
                        "x": float(bounds["X"]),
                        "y": float(bounds["Y"]),
                        "width": win_width,
                        "height": win_height,
                        "is_fullscreen": False,
                    }
                    print(f"[Parser] 检测到PPT普通模式，窗口位置：{ppt_win}")

        # 兜底
        if not ppt_win:
            ppt_win = {
                "x": 0.0,
                "y": 0.0,
                "width": screen_width,
                "height": screen_height,
                "is_fullscreen": True,
            }
            print(
                f"[Parser] 未找到PPT窗口，强制使用屏幕坐标：{screen_width}x{screen_height}"
            )

        self.ppt_win = ppt_win

        # 计算缩放
        scale_x = ppt_win["width"] / sw
        scale_y = ppt_win["height"] / sh
        scale = min(scale_x, scale_y)

        # 解析文本元素
        elem_index = 0
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                full_text = shape.text.strip()

                x = ppt_win["x"] + (float(shape.left) * scale)
                y = ppt_win["y"] + (float(shape.top) * scale) + 100
                w = float(shape.width) * scale
                h = float(shape.height) * scale

                self.elem_texts[elem_index] = full_text
                self.elem_coords[elem_index] = (x, y, w, h)
                print(f"[Parser] 元素{elem_index}：{full_text} → 坐标({x:.2f},{y:.2f})")
                elem_index += 1

    def get_elem_texts(self):
        return self.elem_texts

    def get_elem_coords(self):
        return self.elem_coords

    def list_elems(self):
        res = []
        for idx in self.elem_texts:
            res.append(
                {
                    "index": idx,
                    "text": self.elem_texts[idx],
                    "coords": self.elem_coords[idx],
                }
            )
        return res
